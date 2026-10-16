"""Regenerate files in templates/: starter placeholder template, demo letterhead, demo sample letter.
Run:  python tools/build_assets.py"""
from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT  # noqa: E402
from docx.shared import Cm, Pt  # noqa: E402
from PIL import Image, ImageDraw, ImageFont  # noqa: E402

from lettergen import themes  # noqa: E402
from lettergen.docx_writer import apply_page_setup, bottom_border, set_default_style, set_run_font  # noqa: E402

T = ROOT / "templates"
THEME = themes.resolve("CBIC Classic")


def _font(size, bold=False):
    for name in (["DejaVuSerif-Bold.ttf", "LiberationSerif-Bold.ttf"] if bold else
                 ["DejaVuSerif.ttf", "LiberationSerif-Regular.ttf"]):
        for base in ("/usr/share/fonts/truetype/dejavu", "/usr/share/fonts/truetype/liberation"):
            p = Path(base) / name
            if p.exists():
                return ImageFont.truetype(str(p), size)
    return ImageFont.load_default()


def letterhead_png():
    w, h = 2480, 420
    img = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(img)
    d.rectangle([40, 40, 300, 380], outline="#999999", width=4)
    d.text((170, 210), "LOGO", fill="#999999", font=_font(48, True), anchor="mm")
    lines = [("GOVERNMENT OF INDIA", 58, True), ("MINISTRY OF FINANCE, DEPARTMENT OF REVENUE", 44, False),
             ("OFFICE OF THE COMMISSIONER OF CGST & CENTRAL EXCISE", 52, True),
             ("[YOUR COMMISSIONERATE / DIVISION]", 50, True),
             ("[Address line]  |  Phone: [ ]  |  Email: [ ]", 36, False)]
    y = 60
    for text, size, bold in lines:
        d.text((w // 2 + 130, y), text, fill="black", font=_font(size, bold), anchor="ma")
        y += size + 22
    d.line([(40, h - 12), (w - 40, h - 12)], fill="black", width=6)
    img.save(T / "sample_letterhead.png")


def placeholder_template():
    doc = Document()
    set_default_style(doc, THEME)
    sec = doc.sections[0]
    apply_page_setup(sec, THEME)
    sec.different_first_page_header_footer = True
    hdr = sec.first_page_header.paragraphs[0]
    hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(hdr.add_run("[ Replace this header text with your letterhead image: "
                             "Insert ▸ Header ▸ Edit Header ▸ Insert Picture ]"), THEME, size=10, italic=True)
    bottom_border(hdr)
    width = 21.0 - THEME["margin_left"] - THEME["margin_right"]

    def p(text="", bold=False, align=None, after=6):
        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(after)
        if align:
            para.alignment = align
        if text:
            set_run_font(para.add_run(text), THEME, bold=bold)
        return para

    q = p(after=2)
    q.paragraph_format.tab_stops.add_tab_stop(Cm(width), WD_TAB_ALIGNMENT.RIGHT)
    set_run_font(q.add_run("F.No. {{ FILE_NO }}\tDate: {{ DATE }}"), THEME)
    p("{%p if DIN %}", after=0)
    p("DIN: {{ DIN }}", bold=True, after=2)
    p("{%p endif %}", after=0)
    p("{{p TO_BLOCK }}", after=0)
    q = p(after=6)
    q.paragraph_format.left_indent = Cm(1.6)
    q.paragraph_format.first_line_indent = Cm(-1.6)
    q.paragraph_format.tab_stops.add_tab_stop(Cm(1.6))
    set_run_font(q.add_run("Sub:\t{{ SUBJECT }}"), THEME, bold=True)
    p("{{p REF_BLOCK }}", after=0)
    p("{{ SALUTATION }}")
    p("{{p BODY }}", after=0)
    p("{{p SIGNATURE_BLOCK }}", after=0)
    p("{{p ENCL_BLOCK }}", after=0)
    p("{{p COPY_BLOCK }}", after=0)
    doc.save(T / "CBIC_placeholder_template.docx")


def sample_letter():
    """A plain 'old letter' like one an office would already have - used to demo clone mode."""
    doc = Document()
    th = dict(THEME, font="Bookman Old Style", size=11.5)
    set_default_style(doc, th)
    apply_page_setup(doc.sections[0], th)

    def p(text="", bold=False, align=None, after=4, left=None, first=None, underline=False):
        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(after)
        para.paragraph_format.line_spacing = 1.5
        if align is not None:
            para.alignment = align
        if left is not None:
            para.paragraph_format.left_indent = Cm(left)
        if first is not None:
            para.paragraph_format.first_line_indent = Cm(first)
        if text:
            set_run_font(para.add_run(text), th, bold=bold, underline=underline)
        return para

    lh = doc.add_paragraph()
    lh.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lh.add_run().add_picture(str(T / "sample_letterhead.png"), width=Cm(16.4))
    q = p()
    q.paragraph_format.tab_stops.add_tab_stop(Cm(16.4), WD_TAB_ALIGNMENT.RIGHT)
    set_run_font(q.add_run("C.No. GEXCOM/LGL/OLD/01/2025\tDate: 01.01.2025"), th)
    p("To", after=0)
    p("The Assistant Commissioner,", left=1.0, after=0)
    p("Tirunelveli Division.", left=1.0)
    p("Madam/Sir,")
    q = p(left=1.6, first=-1.6, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    set_run_font(q.add_run("Subject: "), th, bold=True)
    set_run_font(q.add_run("Old sample subject - reg."), th, bold=True, underline=True)
    for i in range(1, 4):
        p(f"{i}. This is an old body paragraph number {i} of the sample letter, which will be removed and "
          "replaced by the new letter while the letterhead, margins, fonts and footer are kept as they are.",
          align=WD_ALIGN_PARAGRAPH.JUSTIFY, first=1.27)
    p("Yours faithfully,", left=10.0)
    p()
    p("(SIGNATORY)", bold=True, left=10.0, after=0)
    p("Assistant Commissioner (Legal)", left=10.0)
    f = doc.sections[0].footer.paragraphs[0]
    f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(f.add_run("Footer of the sample letter (kept)"), th, size=9)
    doc.save(T / "sample_letter_for_clone_demo.docx")


if __name__ == "__main__":
    T.mkdir(exist_ok=True)
    letterhead_png()
    placeholder_template()
    sample_letter()
    print("templates rebuilt in", T)
