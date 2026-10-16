"""Low-level python-docx helpers + the writer that lays out a CBIC-style letter on any container
(a Document, or a docxtpl Subdoc) following a resolved theme."""
from __future__ import annotations

import re

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from .schema import format_date

ROMAN = ["i", "ii", "iii", "iv", "v", "vi", "vii", "viii", "ix", "x", "xi", "xii", "xiii", "xiv", "xv",
         "xvi", "xvii", "xviii", "xix", "xx"]
DEVANAGARI = re.compile(r"[ऀ-ॿ]")


# --------------------------------------------------------------------------- fonts
def set_run_font(run, theme: dict, size=None, bold=None, italic=None, underline=None):
    font = theme.get("font", "Times New Roman")
    hindi = theme.get("hindi_font", "Nirmala UI")
    size = size or theme.get("size", 12)
    run.font.name = font
    run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        rfonts.set(qn(attr), font)
    rfonts.set(qn("w:cs"), hindi)
    szcs = rpr.find(qn("w:szCs"))
    if szcs is None:
        szcs = OxmlElement("w:szCs")
        rpr.append(szcs)
    szcs.set(qn("w:val"), str(int(round(float(size) * 2))))
    if bold is not None:
        run.bold = bold
        if bold:
            b = OxmlElement("w:bCs")
            rpr.append(b)
    if italic is not None:
        run.italic = italic
    if underline is not None:
        run.underline = underline
    return run


def set_default_style(doc, theme: dict):
    st = doc.styles["Normal"]
    st.font.name = theme.get("font", "Times New Roman")
    st.font.size = Pt(theme.get("size", 12))
    rpr = st.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        rfonts.set(qn(attr), theme.get("font", "Times New Roman"))
    rfonts.set(qn("w:cs"), theme.get("hindi_font", "Nirmala UI"))
    pf = st.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(theme.get("space_after", 6))
    pf.line_spacing = theme.get("line_spacing", 1.15)


def text_width_cm(theme: dict, section=None) -> float:
    if section is not None and section.page_width:
        return (section.page_width - section.left_margin - section.right_margin) / 360000
    page = 21.0 if theme.get("page_size", "A4") == "A4" else 21.59
    return page - theme.get("margin_left", 2.54) - theme.get("margin_right", 2.0)


def apply_page_setup(section, theme: dict):
    if theme.get("page_size", "A4") == "A4":
        section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    elif theme.get("page_size") == "Legal":
        section.page_width, section.page_height = Cm(21.59), Cm(35.56)
    else:
        section.page_width, section.page_height = Cm(21.59), Cm(27.94)
    section.top_margin = Cm(theme.get("margin_top", 2.0))
    section.bottom_margin = Cm(theme.get("margin_bottom", 2.0))
    section.left_margin = Cm(theme.get("margin_left", 2.54))
    section.right_margin = Cm(theme.get("margin_right", 2.0))
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)


def bottom_border(paragraph, size=12, color="000000"):
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), str(size))
    b.set(qn("w:space"), "1")
    b.set(qn("w:color"), color)
    pbdr.append(b)
    ppr.append(pbdr)


def add_field(paragraph, instr: str, theme: dict, size=None):
    """Insert a Word field (PAGE, NUMPAGES ...) as begin / instr / separate / result / end runs."""
    def run_with(*children):
        r = paragraph.add_run()
        set_run_font(r, theme, size=size)
        for ch in children:
            r._r.append(ch)
        return r

    def fld(kind):
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), kind)
        return el

    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    t = OxmlElement("w:t")
    t.text = "1"
    run_with(fld("begin"))
    run_with(it)
    run_with(fld("separate"))
    run_with(t)
    run_with(fld("end"))


def set_table_borders(table, sz=4):
    tbl = table._tbl
    tblpr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tblpr.append(borders)


# --------------------------------------------------------------------------- writer
class LetterWriter:
    """Writes letter parts into `container` (anything with add_paragraph / add_table)."""

    def __init__(self, container, theme: dict, width_cm: float):
        self.c = container
        self.t = theme
        self.width = width_cm
        self.labels = theme.get("labels", {})

    # basic paragraph ----------------------------------------------------------
    def para(self, text="", bold=False, italic=False, underline=False, align=None, size=None,
             left=None, first=None, after=None, before=None, keep_next=False, spacing=None):
        p = self.c.add_paragraph()
        pf = p.paragraph_format
        pf.space_after = Pt(self.t.get("space_after", 6) if after is None else after)
        pf.space_before = Pt(before or 0)
        pf.line_spacing = spacing or self.t.get("line_spacing", 1.15)
        if left is not None:
            pf.left_indent = Cm(left)
        if first is not None:
            pf.first_line_indent = Cm(first)
        if align:
            pf.alignment = align
        if keep_next:
            pf.keep_with_next = True
        if text:
            self.runs(p, text, bold=bold, italic=italic, underline=underline, size=size)
        return p

    def runs(self, p, text, bold=False, italic=False, underline=False, size=None):
        """Add text, honouring **bold** markers and line breaks."""
        parts = re.split(r"(\*\*[^*]+\*\*)", str(text))
        for part in parts:
            if not part:
                continue
            b = bold
            if part.startswith("**") and part.endswith("**"):
                part, b = part[2:-2], True
            lines = part.split("\n")
            for j, line in enumerate(lines):
                r = p.add_run(line)
                set_run_font(r, self.t, size=size, bold=b, italic=italic, underline=underline)
                if j < len(lines) - 1:
                    r.add_break()
        return p

    def label_value(self, label, value, bold_label=False, bold_value=False, underline_value=False,
                    hanging=True, after=None, justify=False):
        """'Sub:' <tab> value  with hanging indent so wrapped lines align with the value."""
        ind = self.t.get("label_indent", 1.6 if len(label) <= 6 else 2.6)
        p = self.para(after=after, align=WD_ALIGN_PARAGRAPH.JUSTIFY if justify else None)
        if hanging:
            p.paragraph_format.left_indent = Cm(ind)
            p.paragraph_format.first_line_indent = Cm(-ind)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(ind))
        self.runs(p, label, bold=bold_label)
        self.runs(p, "\t" if hanging else " ")
        self.runs(p, value, bold=bold_value, underline=underline_value)
        return p

    # letter parts -------------------------------------------------------------
    def file_date_block(self, letter):
        L = self.labels
        date = format_date(letter.get("date"), self.t.get("date_style", "%d.%m.%Y"))
        file_no = letter.get("file_no") or ""
        same = self.t.get("file_and_date_same_line", True)
        if same:
            p = self.para(after=2)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(self.width), WD_TAB_ALIGNMENT.RIGHT)
            self.runs(p, f"{L.get('file_no', 'F.No.')} {file_no}".strip(), bold=self.t.get("file_no_bold", False))
            self.runs(p, f"\t{L.get('date', 'Date:')} {date}")
        else:
            self.para(f"{L.get('file_no', 'F.No.')} {file_no}".strip(), after=2)
        if letter.get("din"):
            self.para(f"{L.get('din', 'DIN:')} {letter['din']}", bold=True, after=2)
        if not same:
            self.para(f"{L.get('date', 'Date:')} {date}", align=WD_ALIGN_PARAGRAPH.RIGHT, after=2)
        self.para(after=2)

    def addressees(self, letter):
        L = self.labels
        tos = letter.get("to", [])
        self.para(L.get("to", "To,"), after=0, keep_next=True)
        for n, a in enumerate(tos, 1):
            lines = [x for x in (a.get("name"), a.get("designation"), a.get("organisation")) if x]
            lines += a.get("address_lines", [])
            if a.get("gstin"):
                lines.append(f"GSTIN: {a['gstin']}")
            if a.get("email"):
                lines.append(f"Email: {a['email']}")
            for i, line in enumerate(lines):
                prefix = f"{n}. " if len(tos) > 1 and i == 0 else ""
                self.para(prefix + line, left=0.6 + (0.6 if len(tos) > 1 and i else 0), after=0, keep_next=True)
            self.para(after=2)
        if letter.get("through"):
            self.label_value(L.get("through", "Through:"), letter["through"])
        if letter.get("kind_attention"):
            self.label_value(L.get("kind_attention", "Kind Attention:"), letter["kind_attention"])

    def salutation(self, letter):
        text = self.t.get("salutation_override") or letter.get("salutation") or "Sir/Madam,"
        self.para(text, after=6, keep_next=True)

    def subject(self, letter):
        L = self.labels
        self.label_value(L.get("subject", "Sub:"), letter.get("subject", ""),
                         bold_label=self.t.get("subject_bold", True), bold_value=self.t.get("subject_bold", True),
                         underline_value=self.t.get("subject_underline", False),
                         hanging=self.t.get("subject_hanging", True), justify=True)

    def references(self, letter):
        refs = letter.get("references", [])
        if not refs:
            return
        L = self.labels
        if len(refs) == 1:
            self.label_value(L.get("reference", "Ref:"), refs[0], justify=True)
        else:
            ind = 1.6
            for i, r in enumerate(refs):
                lab = L.get("reference", "Ref:") if i == 0 else ""
                p = self.para(after=2, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
                pf = p.paragraph_format
                pf.left_indent = Cm(ind + 1.0)
                pf.first_line_indent = Cm(-(ind + 1.0))
                pf.tab_stops.add_tab_stop(Cm(ind))
                pf.tab_stops.add_tab_stop(Cm(ind + 1.0))
                self.runs(p, f"{lab}\t{i + 1}.\t{r}")
            p.paragraph_format.space_after = Pt(self.t.get("space_after", 6) + 4)
        if self.t.get("stars_after_ref"):
            self.para("*****", align=WD_ALIGN_PARAGRAPH.CENTER, after=4)

    def body(self, letter):
        mode = self.t.get("para_numbering", "all")
        justify = WD_ALIGN_PARAGRAPH.JUSTIFY if self.t.get("justify", True) else WD_ALIGN_PARAGRAPH.LEFT
        first = self.t.get("first_line_indent", 1.27)
        n = 0
        for item in letter.get("body_paragraphs", []):
            if isinstance(item, dict) and "table" in item:
                self.table(item["table"])
                continue
            text = item if isinstance(item, str) else item.get("text", "")
            points = [] if isinstance(item, str) else item.get("points", [])
            n += 1
            num = ""
            if mode == "all" or (mode == "from_second" and n > 1):
                num = f"{n}.   "
            p = self.para(align=justify, first=first)
            self.runs(p, num + text)
            for j, pt in enumerate(points):
                r = ROMAN[j] if j < len(ROMAN) else str(j + 1)
                q = self.para(align=justify, left=first + 1.0, first=-1.0)
                q.paragraph_format.tab_stops.add_tab_stop(Cm(first + 1.0))
                self.runs(q, f"({r})\t{pt}")

    def table(self, t: dict):
        headers, rows = t.get("headers", []), t.get("rows", [])
        if t.get("caption"):
            self.para(t["caption"], bold=True, after=2, keep_next=True)
        ncols = max([len(headers)] + [len(r) for r in rows] + [1])
        tbl = self.c.add_table(rows=0, cols=ncols)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(tbl)
        size = max(self.t.get("size", 12) - 1, 9)

        def fill(row_vals, bold=False):
            cells = tbl.add_row().cells
            for i in range(ncols):
                val = row_vals[i] if i < len(row_vals) else ""
                para = cells[i].paragraphs[0]
                para.paragraph_format.space_after = Pt(1)
                numeric = bool(re.fullmatch(r"[₹Rs.\s]*[\d,]+(\.\d+)?(/-)?", str(val).strip())) and i > 0
                para.alignment = WD_ALIGN_PARAGRAPH.RIGHT if numeric else WD_ALIGN_PARAGRAPH.LEFT
                self.runs(para, str(val), bold=bold, size=size)

        if headers:
            fill(headers, bold=True)
        for r in rows:
            fill(r, bold=bool(r and re.search(r"total", str(r[0]), re.I)))
        self.para(after=4)

    def closing_and_signature(self, letter):
        sig = letter.get("signatory", {})
        right = self.t.get("signature_align", "right") == "right"
        left = max(self.width - 7.0, 0) if right else 0
        closing = self.t.get("closing_override") or letter.get("closing") or "Yours faithfully,"
        self.para(closing, left=left, before=6, after=0, keep_next=True)
        for _ in range(self.t.get("signature_gap_lines", 3)):
            self.para(after=0, keep_next=True)
        name = sig.get("name", "")
        if name:
            if self.t.get("name_uppercase"):
                name = name.upper()
            if self.t.get("name_in_brackets", True) and not name.startswith("("):
                name = f"({name})"
            self.para(name, bold=True, left=left, after=0, keep_next=True)
        tail = [x for x in (sig.get("designation"), sig.get("office_line")) if x]
        for i, line in enumerate(tail):
            self.para(line, left=left, after=0, keep_next=i < len(tail) - 1)
        self.para(after=4)

    def enclosures(self, letter):
        enc = letter.get("enclosures", [])
        if not enc:
            return
        L = self.labels
        if len(enc) == 1:
            self.label_value(L.get("encl", "Encl.:"), enc[0], hanging=True)
        else:
            self.para(L.get("encl", "Encl.:"), bold=True, after=0, keep_next=True)
            for i, e in enumerate(enc, 1):
                self.para(f"{i}.\t{e}", left=1.6, first=-0.8, after=0)
        self.para(after=2)

    def copy_to(self, letter):
        cc = letter.get("copy_to", [])
        if not cc:
            return
        self.para(self.labels.get("copy_to", "Copy to:"), bold=True, after=2, keep_next=True)
        for i, x in enumerate(cc, 1):
            p = self.para(after=2, left=1.4, first=-0.8, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(1.4))
            self.runs(p, f"{i}.\t{x}")

    def full_letter(self, letter):
        self.file_date_block(letter)
        self.addressees(letter)
        if self.t.get("salutation_position", "after_subject") == "before_subject":
            self.salutation(letter)
        self.subject(letter)
        self.references(letter)
        if self.t.get("salutation_position", "after_subject") != "before_subject":
            self.salutation(letter)
        self.body(letter)
        self.closing_and_signature(letter)
        self.enclosures(letter)
        self.copy_to(letter)
