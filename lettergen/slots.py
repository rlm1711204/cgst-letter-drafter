"""Exact-format engine: fill a Word file that contains [[MARKERS]] while keeping every paragraph,
table, font, spacing and numbering setting of the original.

Marker syntax (type them in Word to make your own format):
  [[KEY]]                 value inserted in place (dotted keys allowed: [[centre.zone]])
  [[?KEY]] ...            paragraph is deleted when KEY is empty
  [[@KEY]]                paragraph repeated once per item of list KEY
  [[%REGION:style]]       prototype paragraph; data[REGION] = [{"style": ..., "text": ...}, ...]
                          is laid out using a copy of the prototype of the same style
  [[IMG:KEY]]             picture from images[KEY] (path) - paragraph stays empty if no picture
Inside values: **bold**, __underline__, a newline gives a line break, a tab character gives a tab.
"""
from __future__ import annotations

import copy
import re

from docx import Document
from lxml import etree
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu
from docx.text.paragraph import Paragraph

W_P, W_R, W_T, W_RPR, W_PPR = qn("w:p"), qn("w:r"), qn("w:t"), qn("w:rPr"), qn("w:pPr")
MARK = re.compile(r"\[\[([^\[\]]+?)\]\]")
MARKUP = re.compile(r"(\*\*|__|\n|\t)")


# --------------------------------------------------------------------------- helpers
def p_text(p) -> str:
    return "".join(t.text or "" for t in p.iter(W_T))


def lookup(data: dict, key: str):
    cur = data
    for part in key.split("."):
        if isinstance(cur, dict):
            cur = cur.get(part, "")
        else:
            return ""
    return cur


def as_text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, (list, tuple)):
        return "\n".join(as_text(x) for x in v if as_text(x))
    if isinstance(v, dict):
        return v.get("text", "") or "; ".join(f"{k}: {as_text(x)}" for k, x in v.items() if as_text(x))
    return str(v)


def _merge_runs(p):
    """Join neighbouring plain-text runs with identical formatting (Word splits [[MARKERS]] otherwise)."""
    runs = [r for r in p if r.tag == W_R]
    prev = None
    for r in runs:
        kids = [k for k in r if k.tag != W_RPR]
        plain = all(k.tag == W_T for k in kids) and kids
        if prev is not None and plain and r.getprevious() is prev:
            a, b = prev.find(W_RPR), r.find(W_RPR)
            sa = b"" if a is None else etree.tostring(a)
            sb = b"" if b is None else etree.tostring(b)
            if sa == sb:
                pt = prev.findall(W_T)[-1]
                pt.text = (pt.text or "") + "".join(t.text or "" for t in r.findall(W_T))
                pt.set(qn("xml:space"), "preserve")
                p.remove(r)
                continue
        prev = r if plain else None


def _runs_for(text: str, base_rpr):
    """Build w:r elements for text with **bold** / __underline__ / newline / tab support."""
    out, bold, under = [], False, False
    for tok in MARKUP.split(text):
        if tok == "":
            continue
        if tok == "**":
            bold = not bold
            continue
        if tok == "__":
            under = not under
            continue
        r = OxmlElement("w:r")
        rpr = copy.deepcopy(base_rpr) if base_rpr is not None else OxmlElement("w:rPr")
        if bold and rpr.find(qn("w:b")) is None:
            b = OxmlElement("w:b")
            rpr.insert(0, b)
        if under and rpr.find(qn("w:u")) is None:
            u = OxmlElement("w:u")
            u.set(qn("w:val"), "single")
            rpr.append(u)
        if len(rpr):
            r.append(rpr)
        if tok == "\n":
            r.append(OxmlElement("w:br"))
        elif tok == "\t":
            r.append(OxmlElement("w:tab"))
        else:
            t = OxmlElement("w:t")
            t.set(qn("xml:space"), "preserve")
            t.text = tok
            r.append(t)
        out.append(r)
    return out


def _base_rpr(p):
    for r in p.iter(W_R):
        if r.find(W_T) is not None:
            return r.find(W_RPR)
    return None


def set_paragraph_text(p, text: str):
    """Replace all runs of paragraph p by `text`, using the formatting of its first text run."""
    base = _base_rpr(p)
    base = copy.deepcopy(base) if base is not None else None
    for child in list(p):
        if child.tag != W_PPR:
            p.remove(child)
    for r in _runs_for(text, base):
        p.append(r)


def _replace_inline(p, data):
    """Replace [[KEY]] markers inside runs, keeping each run's formatting."""
    for r in list(p.iter(W_R)):
        ts = r.findall(W_T)
        if not ts:
            continue
        text = "".join(t.text or "" for t in ts)
        if "[[" not in text:
            continue
        base = r.find(W_RPR)
        parts = MARK.split(text)
        new_runs = []
        for i, part in enumerate(parts):
            if i % 2 == 0:
                if part:
                    new_runs += _runs_for(part, base)
            else:
                new_runs += _runs_for(as_text(lookup(data, part.strip())), base)
        parent = r.getparent()
        idx = parent.index(r)
        parent.remove(r)
        for j, nr in enumerate(new_runs):
            parent.insert(idx + j, nr)


def _strip_leading_marker(p):
    for t in p.iter(W_T):
        if t.text and t.text.startswith("[["):
            t.text = MARK.sub("", t.text, count=1)
            return


# --------------------------------------------------------------------------- numbering
def restart_numbering(doc):
    """Each separate block of auto-numbered items starts again at 1 (Word would continue otherwise)."""
    try:
        numbering = doc.part.numbering_part.element
    except Exception:
        return
    seen, prev_num = set(), None
    for p in doc.element.body.iter(W_P):
        num_pr = p.find(f"{W_PPR}/{qn('w:numPr')}")
        nid = num_pr.find(qn("w:numId")) if num_pr is not None else None
        if nid is None:
            if p_text(p).strip():
                prev_num = None
            continue
        val = nid.get(qn("w:val"))
        if val == prev_num:
            continue
        if val in seen:
            num = numbering.num_having_numId(int(val))
            new = numbering.add_num(num.abstractNumId.val)
            new.add_lvlOverride(ilvl=0).add_startOverride(1)
            new_id = str(new.numId)
            # renumber this whole contiguous block
            q = p
            while q is not None:
                qp = q.find(f"{W_PPR}/{qn('w:numPr')}/{qn('w:numId')}")
                if qp is None:
                    if p_text(q).strip():
                        break
                elif qp.get(qn("w:val")) == val:
                    qp.set(qn("w:val"), new_id)
                q = q.getnext()
            prev_num = new_id
            continue
        seen.add(val)
        prev_num = val


# --------------------------------------------------------------------------- main
def render(template_path, data: dict, images: dict | None = None, regions: dict | None = None):
    """Return a python-docx Document with all markers filled."""
    images = images or {}
    regions = regions or {}
    doc = Document(str(template_path))
    roots = [doc.element.body]
    for s in doc.sections:
        for hf in (s.header, s.footer, s.first_page_header, s.first_page_footer):
            try:
                if not hf.is_linked_to_previous:
                    roots.append(hf._element)
            except Exception:
                pass
    for root in roots:
        for p in list(root.iter(W_P)):
            _merge_runs(p)

    # 1. regions (prototype paragraphs)
    protos: dict[str, dict] = {}
    anchors: dict[str, object] = {}
    for p in list(doc.element.body.iter(W_P)):
        m = re.match(r"\s*\[\[%([A-Za-z0-9_]+):([A-Za-z0-9_]+)\]\]", p_text(p))
        if m:
            region, style = m.groups()
            protos.setdefault(region, {})[style] = p
            anchors.setdefault(region, p)
    for region, styles in protos.items():
        cfg = regions.get(region, {})
        counted = set(cfg.get("count_styles", ["first", "numbered"]))
        prefixes = cfg.get("prefix", {"numbered": "{n}. \t"})
        anchor = anchors[region]
        n = 0
        for block in data.get(region) or []:
            if isinstance(block, str):
                block = {"style": cfg.get("default_style", next(iter(styles))), "text": block}
            style = block.get("style") if block.get("style") in styles else cfg.get("default_style",
                                                                                      next(iter(styles)))
            if style in counted:
                n += 1
            proto = styles[style]
            new = copy.deepcopy(proto)
            text = prefixes.get(style, "").replace("{n}", str(n)) + str(block.get("text", ""))
            set_paragraph_text(new, text)
            anchor.addprevious(new)
        for p in styles.values():
            p.getparent().remove(p)

    # 2. images, conditionals, repeats, inline values
    for root in roots:
        for p in list(root.iter(W_P)):
            text = p_text(p)
            if "[[" not in text:
                continue
            m = re.match(r"\s*\[\[IMG:([A-Za-z0-9_.]+)\]\]", text)
            if m:
                path = images.get(m.group(1))
                for child in list(p):
                    if child.tag != W_PPR:
                        p.remove(child)
                if path:
                    width = (regions.get("_images", {}).get(m.group(1)) or {}).get("width_emu")
                    run = Paragraph(p, doc._body).add_run()
                    run.add_picture(str(path), width=Emu(width) if width else None)
                continue
            m = re.match(r"\s*\[\[\?([A-Za-z0-9_.]+)\]\]", text)
            if m:
                if not as_text(lookup(data, m.group(1))).strip():
                    p.getparent().remove(p)
                    continue
                _strip_leading_marker(p)
                text = p_text(p)
            m = re.match(r"\s*\[\[@([A-Za-z0-9_.]+)\]\]\s*$", text)
            if m:
                items = lookup(data, m.group(1)) or []
                for item in (items if isinstance(items, list) else [items]):
                    new = copy.deepcopy(p)
                    set_paragraph_text(new, as_text(item))
                    p.addprevious(new)
                p.getparent().remove(p)
                continue
            _replace_inline(p, data)
    restart_numbering(doc)
    return doc


def markers(template_path) -> list[str]:
    """List the markers a template expects (for documentation / validation)."""
    doc = Document(str(template_path))
    found = []
    for p in doc.element.body.iter(W_P):
        for m in MARK.finditer(p_text(p)):
            if m.group(1) not in found:
                found.append(m.group(1))
    return found
