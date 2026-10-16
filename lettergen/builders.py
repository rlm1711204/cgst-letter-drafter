"""Three ways to produce the Word file:

1. theme       - build from scratch: letterhead image (or text letterhead) + chosen theme
2. clone       - open YOUR sample letter .docx, keep its letterhead / header / footer / margins /
                 fonts, delete the old letter content and write the new letter in the detected style
3. placeholder - render a .docx that contains {{ TAGS }} (docxtpl / Jinja) - exact layout control
"""
from __future__ import annotations

import collections
import re
import shutil
import subprocess
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from . import __version__
from .docx_writer import (LetterWriter, add_field, apply_page_setup, bottom_border, set_default_style,
                          set_run_font, text_width_cm)
from .schema import format_date

W_P, W_TBL, W_SECT = qn("w:p"), qn("w:tbl"), qn("w:sectPr")

LABEL_PATTERNS = {
    "file_no": r"^\s*((?:फा\.?\s*सं\.?\s*/?\s*)?(?:F|C|File|Letter|O\.?C)\s*\.?\s*No\.?\s*[:.\-]?)",
    "din": r"^\s*((?:डिन\s*/?\s*)?DIN\s*(?:No\.?)?\s*[:.\-]?)",
    "date": r"((?:दिनांक\s*/?\s*)?Dated?\s*[:.\-]?)\s*\d",
    "to": r"^\s*((?:सेवा में\s*/?\s*)?To\s*,?)\s*$",
    "subject": r"^\s*((?:विषय\s*/?\s*)?Sub(?:ject)?\s*\.?\s*[:\-–]+)",
    "reference": r"^\s*((?:संदर्भ\s*/?\s*)?Ref(?:erence)?s?\s*\.?\s*[:\-–]+)",
    "copy_to": r"^\s*((?:प्रतिलिपि\s*/?\s*)?Copy\s*(?:to)?\s*[:.\-]?)",
    "encl": r"^\s*((?:संलग्नक?\s*/?\s*)?Encl(?:osures?)?\s*\.?\s*[:.\-]?)",
    "through": r"^\s*(Through\s*[:\-–]?)",
    "kind_attention": r"^\s*(Kind\s+Attn\.?|Kind\s+Attention\s*[:\-–]?)",
}
ANCHOR_KEYS = ("file_no", "din", "to", "subject")
HEAD_KEYS = ("file_no", "din", "date", "to")
SALUTATION_RE = re.compile(r"^\s*(Respected\s+)?(Dear\s+)?(Sir|Madam|Sir\s*/\s*Madam|Sirs|महोदय)[^.]{0,25}$", re.I)
CLOSING_RE = re.compile(r"^\s*(Yours\s+(faithfully|sincerely|truly)|भवदीय)", re.I)


# =========================================================================== helpers
def _el_text(el) -> str:
    return "".join(t.text or "" for t in el.iter(qn("w:t")))


def _has_image(el) -> bool:
    return any(True for _ in el.iter(qn("w:drawing"))) or any(True for _ in el.iter(qn("w:pict")))


def _xml_parts_text(path) -> str:
    with zipfile.ZipFile(path) as z:
        return " ".join(z.read(n).decode("utf8", "ignore") for n in z.namelist()
                        if re.match(r"word/(document|header\d*|footer\d*)\.xml", n))


def has_placeholders(path) -> bool:
    xml = _xml_parts_text(path)
    plain = re.sub(r"<[^>]+>", "", xml)
    return bool(re.search(r"\{\{.*?\}\}|\{%.*?%\}", plain))


def _defaults(doc):
    font, size = None, None
    st = doc.styles["Normal"]
    if st.font.name:
        font = st.font.name
    if st.font.size:
        size = st.font.size.pt
    dd = doc.styles.element.find(qn("w:docDefaults"))
    if dd is not None:
        rpr = dd.find(f"{qn('w:rPrDefault')}/{qn('w:rPr')}")
        if rpr is not None:
            rf = rpr.find(qn("w:rFonts"))
            if font is None and rf is not None:
                font = rf.get(qn("w:ascii")) or ("Calibri" if rf.get(qn("w:asciiTheme")) else None)
            sz = rpr.find(qn("w:sz"))
            if size is None and sz is not None:
                size = int(sz.get(qn("w:val"))) / 2
    return font or "Calibri", size or 11


def _run_font(run, para, defaults):
    name = run.font.name or (para.style.font.name if para.style is not None else None) or defaults[0]
    size = run.font.size.pt if run.font.size else (
        para.style.font.size.pt if para.style is not None and para.style.font.size else defaults[1])
    return name, size


# =========================================================================== analyse sample
def analyze_sample(path) -> dict:
    """Inspect a sample letter and return what will be kept and the style detected."""
    doc = Document(str(path))
    body = doc.element.body
    elements = [el for el in body if el.tag != W_SECT]
    rep = {"path": str(path), "placeholders": has_placeholders(path), "elements": [], "anchor_index": None,
           "header_images": 0, "header_text": "", "detected": {}, "labels": {}, "notes": []}
    for s in doc.sections:
        for h in (s.header, s.first_page_header):
            if h is not None and not h.is_linked_to_previous:
                rep["header_images"] += sum(1 for _ in h._element.iter(qn("w:drawing")))
                rep["header_text"] += " ".join(p.text for p in h.paragraphs if p.text.strip())[:200]
    for i, el in enumerate(elements):
        txt = _el_text(el).strip()
        kind = "table" if el.tag == W_TBL else "paragraph"
        rep["elements"].append({"index": i, "kind": kind, "image": _has_image(el), "text": txt[:90]})
        if rep["anchor_index"] is None and txt:
            first_line = txt.split("\n")[0]
            if any(re.search(LABEL_PATTERNS[k], first_line, re.I) for k in ANCHOR_KEYS):
                rep["anchor_index"] = i
    if rep["anchor_index"] is None:
        rep["anchor_index"] = 0
        rep["notes"].append("Could not find where the letter starts (File No./DIN/To/Sub). Only the page "
                            "header/footer will be kept. Set KEEP_TOP_BLOCKS manually if the letterhead is in "
                            "the body.")

    # ---- labels & layout ------------------------------------------------------
    paras = doc.paragraphs
    labels, det = {}, {}
    subj_idx = sal_idx = close_idx = None
    for i, p in enumerate(paras):
        t = p.text.strip()
        if not t:
            continue
        for key, pat in LABEL_PATTERNS.items():
            if key in labels or (key in HEAD_KEYS and i > 25):
                continue
            m = re.search(pat, t, 0 if key == "date" else re.I)
            if m:
                labels[key] = re.sub(r"\s+", " ", m.group(1)).strip()
                if key == "file_no":
                    det["file_and_date_same_line"] = bool(re.search(LABEL_PATTERNS["date"], t, re.I))
                if key == "subject":
                    subj_idx = i
                    runs = [r for r in p.runs if r.text.strip()]
                    det["subject_bold"] = any(r.bold for r in runs)
                    det["subject_underline"] = any(r.underline for r in runs[1:]) if len(runs) > 1 else False
        if sal_idx is None and SALUTATION_RE.match(t):
            sal_idx = i
            det["sample_salutation"] = t
        if close_idx is None and CLOSING_RE.match(t):
            close_idx = i
            labels["_closing"] = t
            pf = p.paragraph_format
            indent = (pf.left_indent or 0) / 360000 if pf.left_indent else 0
            det["signature_align"] = "right" if (indent > 4 or p.alignment in (WD_ALIGN_PARAGRAPH.RIGHT,
                                                                                 WD_ALIGN_PARAGRAPH.CENTER)) else "left"
    if subj_idx is not None and sal_idx is not None:
        det["salutation_position"] = "before_subject" if sal_idx < subj_idx else "after_subject"
    if "date" in labels and "file_no" not in labels:
        det["file_and_date_same_line"] = False

    # ---- body style -------------------------------------------------------------
    start = max(x for x in (subj_idx, sal_idx, 0) if x is not None) + 1
    end = close_idx if close_idx is not None else len(paras)
    body_paras = [p for p in paras[start:end] if len(p.text.strip()) > 60]
    defaults = _defaults(doc)
    fonts, sizes = collections.Counter(), collections.Counter()
    for p in (body_paras or paras):
        for r in p.runs:
            if r.text.strip():
                n, s = _run_font(r, p, defaults)
                fonts[n] += len(r.text)
                sizes[s] += len(r.text)
    if fonts:
        det["font"] = fonts.most_common(1)[0][0]
    if sizes:
        det["size"] = sizes.most_common(1)[0][0]
    if body_paras:
        just = sum(1 for p in body_paras if (p.alignment or p.style.paragraph_format.alignment) ==
                   WD_ALIGN_PARAGRAPH.JUSTIFY)
        det["justify"] = just >= len(body_paras) / 2
        fli = [((p.paragraph_format.first_line_indent or 0) / 360000) for p in body_paras]
        det["first_line_indent"] = round(sorted(fli)[len(fli) // 2], 2)
        ls = [p.paragraph_format.line_spacing for p in body_paras if isinstance(p.paragraph_format.line_spacing, float)]
        if ls:
            det["line_spacing"] = collections.Counter(ls).most_common(1)[0][0]
        nums = [bool(re.match(r"^\s*\d{1,2}\.", p.text)) for p in body_paras]
        if nums and nums[0]:
            det["para_numbering"] = "all"
        elif any(nums):
            det["para_numbering"] = "from_second"
        else:
            det["para_numbering"] = "none"
    s0 = doc.sections[0]
    rep["margins_cm"] = {k: round((getattr(s0, f"{k}_margin") or 0) / 360000, 2)
                         for k in ("top", "bottom", "left", "right")}
    rep["labels"] = {k: v for k, v in labels.items() if not k.startswith("_")}
    det["labels"] = dict(rep["labels"])
    if "_closing" in labels:
        det["closing_from_sample"] = labels["_closing"]
    rep["detected"] = det
    return rep


# =========================================================================== builders
def _letterhead_header(section, theme, profile, image_path=None):
    first_only = theme.get("letterhead_on", "first page") == "first page"
    section.different_first_page_header_footer = first_only
    hdr = section.first_page_header if first_only else section.header
    width = text_width_cm(theme, section)
    p = hdr.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if image_path:
        p.add_run().add_picture(str(image_path), width=Cm(width))
        return
    lines = profile.get("office_lines") or []
    for i, line in enumerate(lines):
        q = p if i == 0 else hdr.add_paragraph()
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        q.paragraph_format.space_after = Pt(0)
        set_run_font(q.add_run(line), theme, size=theme.get("size", 12) + (1 if i >= len(lines) - 2 else 0),
                     bold=True)
    for extra in (profile.get("office_address"), profile.get("office_contact")):
        if extra:
            q = hdr.add_paragraph()
            q.alignment = WD_ALIGN_PARAGRAPH.CENTER
            q.paragraph_format.space_after = Pt(0)
            set_run_font(q.add_run(extra), theme, size=max(theme.get("size", 12) - 2, 9))
    bottom_border(hdr.paragraphs[-1])


def _page_numbers(section, theme):
    footers = [section.footer]
    if section.different_first_page_header_footer:
        footers.append(section.first_page_footer)
    for f in footers:
        p = f.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        size = max(theme.get("size", 12) - 3, 8)
        set_run_font(p.add_run("Page "), theme, size=size)
        add_field(p, "PAGE", theme, size=size)
        set_run_font(p.add_run(" of "), theme, size=size)
        add_field(p, "NUMPAGES", theme, size=size)


def _props(doc, letter):
    cp = doc.core_properties
    cp.title = (letter.get("subject") or "")[:250]
    cp.author = letter.get("signatory", {}).get("name") or "CGST Letter Drafter"
    cp.comments = f"Generated by CGST Letter Drafter v{__version__}. Verify before issue."
    cp.keywords = ", ".join(x for x in (letter.get("file_no"), letter.get("din"), letter.get("letter_type")) if x)


def build_theme(letter, theme, profile, letterhead_image=None):
    doc = Document()
    set_default_style(doc, theme)
    sec = doc.sections[0]
    apply_page_setup(sec, theme)
    _letterhead_header(sec, theme, profile, letterhead_image)
    if theme.get("page_numbers", True):
        _page_numbers(sec, theme)
    LetterWriter(doc, theme, text_width_cm(theme, sec)).full_letter(letter)
    _props(doc, letter)
    return doc, {"mode": "theme", "letterhead": "image" if letterhead_image else "text (office profile)"}


def build_clone(letter, theme, sample_path, keep_top=None):
    analysis = analyze_sample(sample_path)
    doc = Document(str(sample_path))
    body = doc.element.body
    keep = analysis["anchor_index"] if keep_top is None or keep_top < 0 else keep_top
    elements = [el for el in body if el.tag != W_SECT]
    for el in elements[keep:]:
        body.remove(el)
    sec = doc.sections[-1]
    footer_has_text = any(p.text.strip() for p in sec.footer.paragraphs) or \
        any(True for _ in sec.footer._element.iter(qn("w:fldChar"))) or \
        any(True for _ in sec.footer._element.iter(qn("w:fldSimple")))
    if theme.get("page_numbers") and not footer_has_text and theme.get("clone_add_page_numbers", False):
        _page_numbers(sec, theme)
    LetterWriter(doc, theme, text_width_cm(theme, sec)).full_letter(letter)
    _props(doc, letter)
    return doc, {"mode": "clone", "kept_blocks": keep, "analysis": analysis}


def placeholder_context(letter, theme, tpl):
    """Context for docxtpl templates. Simple strings + ready-formatted sub-documents."""
    sec = tpl.get_docx().sections[-1]
    width = text_width_cm(theme, sec)

    def sub(fn):
        sd = tpl.new_subdoc()
        w = LetterWriter(sd, theme, width)
        fn(w)
        return sd

    sig = letter.get("signatory", {})
    to_lines = []
    for a in letter.get("to", []):
        to_lines += [x for x in (a.get("name"), a.get("designation"), a.get("organisation")) if x]
        to_lines += a.get("address_lines", [])
        if a.get("gstin"):
            to_lines.append(f"GSTIN: {a['gstin']}")
    ctx = {
        "FILE_NO": letter.get("file_no", ""), "DIN": letter.get("din", ""),
        "DATE": format_date(letter.get("date"), theme.get("date_style", "%d.%m.%Y")),
        "SUBJECT": letter.get("subject", ""), "SALUTATION": letter.get("salutation", ""),
        "CLOSING": letter.get("closing", ""), "THROUGH": letter.get("through", ""),
        "KIND_ATTENTION": letter.get("kind_attention", ""),
        "SIGNATORY_NAME": sig.get("name", ""), "SIGNATORY_DESIGNATION": sig.get("designation", ""),
        "SIGNATORY_OFFICE": sig.get("office_line", ""),
        "TO_LINES": to_lines, "REFERENCES": letter.get("references", []),
        "ENCLOSURES": letter.get("enclosures", []), "COPY_TO": letter.get("copy_to", []),
        "BODY": sub(lambda w: w.body(letter)),
        "TO_BLOCK": sub(lambda w: w.addressees(letter)),
        "REF_BLOCK": sub(lambda w: w.references(letter)),
        "SIGNATURE_BLOCK": sub(lambda w: w.closing_and_signature(letter)),
        "ENCL_BLOCK": sub(lambda w: w.enclosures(letter)),
        "COPY_BLOCK": sub(lambda w: w.copy_to(letter)),
        "FULL_LETTER": sub(lambda w: w.full_letter(letter)),
        "letter": letter, "kf": letter.get("key_facts", {}),
    }
    return ctx


def build_placeholder(letter, theme, template_path):
    from docxtpl import DocxTemplate
    tpl = DocxTemplate(str(template_path))
    ctx = placeholder_context(letter, theme, tpl)
    try:
        missing = tpl.get_undeclared_template_variables(context=ctx)
    except Exception:
        missing = set()
    tpl.render(ctx, autoescape=True)
    _props(tpl.docx, letter)
    return tpl, {"mode": "placeholder", "template": str(template_path), "unknown_tags": sorted(missing or [])}


def generate(letter, theme, profile, mode="auto", template_path=None, letterhead_image=None, keep_top=None):
    """Pick the right builder. Returns (doc_with_save_method, report)."""
    if mode == "auto":
        if template_path and has_placeholders(template_path):
            mode = "placeholder"
        elif template_path:
            mode = "clone"
        else:
            mode = "theme"
    if mode == "placeholder":
        return build_placeholder(letter, theme, template_path)
    if mode == "clone":
        return build_clone(letter, theme, template_path, keep_top)
    return build_theme(letter, theme, profile, letterhead_image)


# =========================================================================== output helpers
def docx_text(path) -> str:
    """Plain text of a .docx in reading order (body paragraphs + tables), for verification prompts."""
    doc = Document(str(path))
    out = []
    for el in doc.element.body:
        if el.tag == W_P:
            out.append(_el_text(el))
        elif el.tag == W_TBL:
            for tr in el.iter(qn("w:tr")):
                out.append(" | ".join(_el_text(tc).strip() for tc in tr.iter(qn("w:tc"))))
    text = "\n".join(out)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def to_pdf(docx_path) -> Path | None:
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        return None
    out_dir = Path(docx_path).parent
    subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", str(out_dir), str(docx_path)],
                   check=False, capture_output=True, timeout=180)
    pdf = Path(docx_path).with_suffix(".pdf")
    return pdf if pdf.exists() else None
