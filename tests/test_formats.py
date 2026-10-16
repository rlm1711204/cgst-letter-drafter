import datetime as dt
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from docx import Document  # noqa: E402

from lettergen import builders, formats, legal, parser, slots  # noqa: E402

TODAY = dt.date(2026, 10, 16)
LH = ROOT / "templates" / "sample_letterhead.png"


def load(name):
    raw, _ = parser.extract_json((ROOT / "examples" / f"{name}.json").read_text(encoding="utf-8"))
    fmt = formats.detect(raw)
    data, _ = formats.normalise(fmt, raw, {})
    return fmt, data


def test_detect_formats():
    assert formats.detect({"format": "sop_checklist"}) == "sop_checklist"
    assert formats.detect({"modus_operandi": "x", "registration": "y"}) == "sop_checklist"
    assert formats.detect({"notesheet_body": []}) == "letter_notesheet"
    assert formats.detect({"subject": "x"}) == "letter"


def test_templates_have_no_case_data_or_images():
    for f in (ROOT / "templates" / "formats").glob("*.docx"):
        import zipfile
        z = zipfile.ZipFile(f)
        assert not [n for n in z.namelist() if n.startswith("word/media/")], f
        import re
        text = re.sub(r"<[^>]+>", "", z.read("word/document.xml").decode("utf8"))
        assert not re.search(r"\d{2}[A-Z]{5}\d{4}[A-Z][0-9A-Z]{3}", text), f   # no GSTIN
        assert "M/s" not in text, f                                            # no party names


def test_sop_fill_all_slots(tmp_path):
    fmt, d = load("sop_checklist")
    doc, _ = formats.build(fmt, d)
    out = tmp_path / "sop.docx"
    doc.save(str(out))
    text = builders.docx_text(out)
    assert "[[" not in text
    assert "33AAKCM4417R1ZZ" in text and "Ineligible ITC" in text
    t = Document(str(out)).tables[0]
    assert len(t.rows) == 26  # exact table structure kept


def test_letter_notesheet_numbering_and_letterhead(tmp_path):
    fmt, d = load("letter_notesheet")
    doc, rep = formats.build(fmt, d, str(LH))
    out = tmp_path / "ln.docx"
    doc.save(str(out))
    text = builders.docx_text(out)
    assert "[[" not in text
    assert "2. In order to verify" in text and "3. Failure" in text
    assert "C.No. GEXCOM/INV/DEMO/01/2026Date: 16.10.2026" in text
    assert "DIN: CBIC-DIN" in text
    d2 = Document(str(out))
    assert d2.inline_shapes and rep["letterhead"] == "image"
    # the two numbered lists use different numIds (notesheet list restarts at 1)
    from docx.oxml.ns import qn
    ids = []
    for p in d2.element.body.iter(qn("w:p")):
        n = p.find(f"{qn('w:pPr')}/{qn('w:numPr')}/{qn('w:numId')}")
        if n is not None and n.get(qn("w:val")) not in ids:
            ids.append(n.get(qn("w:val")))
    assert len(ids) == 2


def test_conditional_din_line_removed(tmp_path):
    fmt, d = load("letter_notesheet")
    d["din"] = ""
    doc, _ = formats.build(fmt, d)
    out = tmp_path / "x.docx"
    doc.save(str(out))
    assert "DIN:" not in builders.docx_text(out)
    issues = formats.run_checks(fmt, d, {}, today=TODAY)
    assert any(i["level"] == "ERROR" and "DIN" in i["msg"] for i in issues)


def test_sop_checks():
    fmt, d = load("sop_checklist")
    d["registration"] = "Registered"
    d["jurisdiction"] = "Centre"
    d["centre"] = {"zone": "", "commissionerate": "", "division": "", "range": ""}
    msgs = [(i["level"], i["msg"]) for i in formats.run_checks(fmt, d, {}, today=TODAY)]
    assert any(lv == "ERROR" and "no GSTIN" in m for lv, m in msgs)
    assert any("Range are blank" in m for lv, m in msgs)


def test_slots_markup_and_restart(tmp_path):
    from docx import Document as D
    doc = D()
    doc.add_paragraph("Hello [[name]], period [[p]]")
    doc.add_paragraph("[[?opt]]Optional [[opt]]")
    doc.add_paragraph("[[@items]]")
    path = tmp_path / "t.docx"
    doc.save(str(path))
    out = slots.render(path, {"name": "**Ram**", "p": "A\nB", "opt": "", "items": ["one", "two"]})
    texts = [slots.p_text(p) for p in out.element.body.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p")]
    assert texts[0] == "Hello Ram, period AB"
    assert "Optional" not in " ".join(texts)
    assert texts[1:3] == ["one", "two"]
    assert out.paragraphs[0].runs[1].bold


def test_time_limits_and_fy_ranges():
    assert legal.financial_years("from FY 2021-22 to FY 2023-24") == [2021, 2022, 2023]
    rows = {(r["fy"], r.get("section")): r for r in legal.demand_time_limits([2022], TODAY)}
    r73 = rows[("2022-23", "73")]
    assert r73["scn_last_date"] == dt.date(2026, 9, 30) and r73["scn_expired"]
    assert not rows[("2022-23", "74")]["scn_expired"]
